"""Deliverable writers: raw pull CSV, scoring intermediate CSV, top-10 Excel.

These map 1:1 onto the brief's supporting deliverables. The Excel column order
is taken verbatim from the mandate config `output.columns`, with sub_score_N
mapped to the configured sub-scores in declaration order.
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from .scoring import ScoredProfile, sg_tenure_years


def _linkedin(p: dict) -> str:
    return p.get("linkedin_url") or p.get("url") or p.get("canonical_url") or ""


def _current_title(p: dict) -> str:
    return p.get("current_title") or p.get("title") or ""


def _current_company(p: dict) -> str:
    return (
        p.get("current_company")
        or p.get("active_experience_company_name")
        or ""
    )


def write_raw_pull(profiles: list[dict], out_path: str | Path) -> Path:
    """The audited filter-precision deliverable: profiles before any ranking."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    cols = [
        "coresignal_id", "name", "linkedin_url", "current_title",
        "current_company", "location_country", "summary_chars",
        "n_experience_roles",
    ]
    with out_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for p in profiles:
            w.writerow({
                "coresignal_id": p.get("id", ""),
                "name": p.get("name", ""),
                "linkedin_url": _linkedin(p),
                "current_title": _current_title(p),
                "current_company": _current_company(p),
                "location_country": p.get("location_country", ""),
                "summary_chars": len(str(p.get("summary", ""))),
                "n_experience_roles": len(p.get("experience", [])),
            })
    return out_path


def write_scoring_intermediate(
    scored: list[ScoredProfile], cfg: dict, ledger, out_path: str | Path
) -> Path:
    """Raw pull + every computed score, before the top-10 cut. Fully auditable:
    deterministic, llm, and blended value for each sub-score, plus rule hits."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    sub_names = list(cfg["scoring"]["sub_scores"].keys())
    cols = ["coresignal_id", "name", "current_title", "current_company",
            "years_relevant_experience", "fit_score", "confidence"]
    for n in sub_names:
        cols += [f"{n}__det", f"{n}__llm", f"{n}__blended", f"{n}__hits"]
    cols += ["credits_spent_on_this_candidate", "flags"]

    with out_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for sp in scored:
            p = sp.profile
            row = {
                "coresignal_id": p.get("id", ""),
                "name": p.get("name", ""),
                "current_title": _current_title(p),
                "current_company": _current_company(p),
                "years_relevant_experience": sp.years_relevant,
                "fit_score": sp.fit_score,
                "confidence": sp.confidence,
                "credits_spent_on_this_candidate": ledger.spent_on_profile(str(p.get("id", ""))),
                "flags": "; ".join(sp.flags),
            }
            for n in sub_names:
                ss = sp.sub_scores[n]
                row[f"{n}__det"] = ss.deterministic
                row[f"{n}__llm"] = ss.llm
                row[f"{n}__blended"] = ss.blended
                row[f"{n}__hits"] = " | ".join(ss.hits)
            w.writerow(row)
    return out_path


def write_top10_excel(
    ranked: list[ScoredProfile], cfg: dict, ledger, out_path: str | Path
) -> Path:
    """Top-N shortlist with the exact deliverable columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    top_n = cfg["ranking"].get("top_n", 10)
    sub_names = list(cfg["scoring"]["sub_scores"].keys())
    columns = cfg["output"]["columns"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Top10 Shortlist"

    # header with the configured sub-score labels appended for readability
    header_friendly = []
    sub_idx = 0
    sub_slots = ["sub_score_1", "sub_score_2", "sub_score_3", "sub_score_4", "sub_score_other"]
    for c in columns:
        if c in sub_slots and sub_idx < len(sub_names):
            label = cfg["scoring"]["sub_scores"][sub_names[sub_idx]]["label"]
            header_friendly.append(f"{c} ({label})")
            sub_idx += 1
        else:
            header_friendly.append(c)
    ws.append(header_friendly)
    head_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for i, sp in enumerate(ranked[:top_n], start=1):
        p = sp.profile
        sub_vals = [sp.sub_scores[n].blended for n in sub_names]
        # pad to 5 sub-slots
        sub_vals += [""] * (5 - len(sub_vals))
        record = {
            "rank": i,
            "coresignal_id": p.get("id", ""),
            "linkedin_url": _linkedin(p),
            "name": p.get("name", ""),
            "current_title": _current_title(p),
            "current_company": _current_company(p),
            "years_relevant_experience": sp.years_relevant,
            "fit_score": sp.fit_score,
            "sub_score_1": sub_vals[0],
            "sub_score_2": sub_vals[1],
            "sub_score_3": sub_vals[2],
            "sub_score_4": sub_vals[3],
            "sub_score_other": sub_vals[4],
            "rationale": sp.rationale,
            "confidence": sp.confidence,
            "credits_spent_on_this_candidate": ledger.spent_on_profile(str(p.get("id", ""))),
            "concerns_or_flags": "; ".join(sp.flags) or "-",
        }
        ws.append([record.get(c, "") for c in columns])

    # column widths
    widths = {"rationale": 60, "concerns_or_flags": 38, "name": 22,
              "current_title": 26, "current_company": 24, "linkedin_url": 30}
    for idx, c in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = widths.get(c, 16)
    ws.freeze_panes = "A2"

    # a second sheet: methodology / weights, for transparency
    ws2 = wb.create_sheet("Scoring Weights")
    ws2.append(["sub_score", "label", "weight", "llm_blend"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    blend_default = cfg["scoring"].get("llm_blend_default")
    for n in sub_names:
        s = cfg["scoring"]["sub_scores"][n]
        ws2.append([n, s["label"], s["weight"], s.get("llm_blend", blend_default)])
    ws2.append([])
    ws2.append(["fit_score = Σ (weight × blended_sub_score)"])
    ws2.append([f"credits spent (total): {ledger.total_spent}"])

    wb.save(out_path)
    return out_path
